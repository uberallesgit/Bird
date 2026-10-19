# -*- coding: utf-8 -*-
import re
import os
import json
import traceback
from datetime import datetime
from pathlib import Path

import openpyxl

from kivy.utils import platform  # <--- ДОБАВЛЕН ИМПОРТ ДЛЯ ОПРЕДЕЛЕНИЯ ОС
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.spinner import Spinner
from kivy.uix.popup import Popup
from kivy.uix.filechooser import FileChooserListView
from kivy.uix.scrollview import ScrollView
from kivy.core.window import Window

# Для Android лучше не задавать жесткий размер окна,
# так как экраны телефонов разные.
if platform not in ('android', 'ios'):
    Window.size = (950, 750)

# ====================================================================
#                      ЛОГИКА И КОНСТАНТЫ
# ====================================================================

PRIORITY_ORDER = ["P0", "P1", "P2", "P3", "НП"]
CONFIG_FILE = "config.json"

DEFAULT_CONSTANTS = {
    "CR1709": "недоступна в связи с прилётом БПЛА",
    "CR0336": "неисправное ТП (БС ушла из статистики)",
}


def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return DEFAULT_CONSTANTS.copy()


def save_config(constants):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(constants, f, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"Не удалось сохранить конфигурацию: {e}")


def parse_priority(val):
    if val is None: return "НП"
    s = str(val).strip()
    if not s: return "НП"
    m = re.search(r"приоритет\s*(\d)", s, re.IGNORECASE)
    return f"P{m.group(1)}" if m else "НП"


def format_duration(delta):
    total = int(delta.total_seconds())
    if total < 0: total = 0
    days = total // 86400
    hours = (total % 86400) // 3600
    minutes = (total % 3600) // 60
    if days >= 1: return f"{days} сут {hours} ч"
    if hours >= 1: return f"{hours} ч {minutes} мин"
    return f"{minutes} мин"


def parse_file_dt(path):
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})_(\d{2})_(\d{2})_(\d{2})", Path(path).name)
    return datetime(*map(int, m.groups())) if m else datetime.now()


def to_dt(val):
    if isinstance(val, datetime): return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return datetime.strptime(val, fmt)
            except ValueError:
                pass
    return None


def label(prio):
    if prio in ("P0", "P1"): return f"**{prio}:**"
    if prio == "P2": return f"*{prio}:*"
    return f"{prio}:"


def label_n(prio, n):
    if prio in ("P0", "P1"): return f"**{prio} ({n} БС):**"
    if prio == "P2": return f"*{prio} ({n} БС):*"
    return f"{prio} ({n} БС):"


def _load_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in wb.sheetnames: return []
        ws = wb[sheet_name]
        return [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    finally:
        wb.close()


def read_om(path):
    return [
        {"alarm": r[0], "occurrence": r[1], "bts": r[2], "sc": r[4], "diesel_priority": r[7] if len(r) > 7 else None}
        for r in _load_sheet(path, "OM")]


def read_diesel(path):
    return [{"bts": r[0], "alarm": r[1], "occurrence": r[2], "sc": r[4] if len(r) > 4 else None} for r in
            _load_sheet(path, "Diesel On")]


def read_alarms(path):
    return [{"bts": r[0], "sc": r[1], "alarm": r[3], "priority": r[5] if len(r) > 5 else None} for r in
            _load_sheet(path, "Sheet")]


def get_unique_scs(om_path):
    try:
        om = read_om(om_path)
        scs = sorted(list(set(r["sc"] for r in om if r["sc"])))
        return scs if scs else ["Сервисный Центр Феодосия"]
    except Exception:
        return ["Сервисный Центр Феодосия"]


def build_ne(om, sc_name, get_prio, constants):
    result, seen = [], set()
    for r in om:
        if r["alarm"] != "NE Is Disconnected" or r["sc"] != sc_name: continue
        bts = r["bts"]
        if bts in seen: continue
        seen.add(bts)
        result.append({"bts": bts, "priority": get_prio(bts), "reason": "причина не указана"})

    for bts, reason in constants.items():
        if bts in seen:
            for x in result:
                if x["bts"] == bts: x["reason"] = reason
        else:
            result.append({"bts": bts, "priority": get_prio(bts), "reason": reason})
            seen.add(bts)
    return result


def build_gen(diesel, sc_name, get_prio, file_dt):
    result, seen = [], set()
    for r in diesel:
        if r["sc"] != sc_name: continue
        bts = r["bts"]
        if bts in seen: continue
        seen.add(bts)
        occ = to_dt(r["occurrence"])
        dur = format_duration(file_dt - occ) if occ else "—"
        result.append({"bts": bts, "priority": get_prio(bts), "duration": dur})
    return result


def build_akb(alarms, sc_name, get_prio, gen_names):
    result, seen = [], set()
    for r in alarms:
        if r["alarm"] != "MAINS FAILURE" or r["sc"] != sc_name: continue
        bts = r["bts"]
        if bts in seen or bts in gen_names: continue
        seen.add(bts)
        result.append({"bts": bts, "priority": get_prio(bts)})
    return result


def _group(items):
    by_p = {}
    for it in items: by_p.setdefault(it["priority"], []).append(it)
    return by_p


def format_ne(items):
    lines, by_p = [], _group(items)
    for p in PRIORITY_ORDER:
        if p not in by_p: continue
        for it in by_p[p]: lines.append(f"- {label(p)} {it['bts']} — {it['reason']}")
    return "\n".join(lines)


def format_akb(items):
    lines, by_p = [], _group(items)
    for p in PRIORITY_ORDER:
        if p not in by_p: continue
        group = by_p[p]
        if len(group) == 1:
            lines.append(f"- {label(p)} {group[0]['bts']}")
        else:
            lines.append(f"- {label_n(p, len(group))} {', '.join(it['bts'] for it in group)}")
    return "\n".join(lines)


def format_gen(items):
    lines, by_p = [], _group(items)
    for p in PRIORITY_ORDER:
        if p not in by_p: continue
        group = by_p[p]
        if len(group) == 1:
            lines.append(f"- {label(p)} {group[0]['bts']} — {group[0]['duration']}")
        else:
            lines.append(f"- {label_n(p, len(group))} {'; '.join(f'{it['bts']} — {it['duration']}' for it in group)}")
    return "\n".join(lines)


def build_report(om_path, alarms_path, sc_name, constants):
    om = read_om(om_path)
    diesel = read_diesel(om_path)
    alarms = read_alarms(alarms_path)
    file_dt = parse_file_dt(om_path)

    om_prio = {r["bts"]: parse_priority(r["diesel_priority"]) for r in om if r["diesel_priority"]}
    alm_prio = {r["bts"]: parse_priority(r["priority"]) for r in alarms if r["sc"] == sc_name and r["priority"]}

    def get_prio(bts): return om_prio.get(bts) or alm_prio.get(bts) or "НП"

    ne = build_ne(om, sc_name, get_prio, constants)
    gen = build_gen(diesel, sc_name, get_prio, file_dt)
    akb = build_akb(alarms, sc_name, get_prio, {g["bts"] for g in gen})

    now = datetime.now()
    out = [
        f"📡 ЦТЭиСО {sc_name.replace('Сервисный Центр ', '')} — сводка", "",
        f"📅 Дата формирования: {now.strftime('%d.%m.%Y')}",
        f"🕐 Время формирования: {now.strftime('%H:%M')}", "", "---", "",
        f"## 🔴 NE DISCONNECTED ({len(ne)} БС)", "", format_ne(ne) if ne else "0 БС", "",
        f"## 🔋 На АКБ: {len(akb)} БС", "", format_akb(akb) if akb else "0 БС", "",
        f"## ⚡ На Генерации: {len(gen)} БС", "", format_gen(gen) if gen else "0 БС"
    ]
    return "\n".join(out)


# ====================================================================
#                        KIVY - ИНТЕРФЕЙС
# ====================================================================

class ReportApp(App):
    title = "Сводка ЦТЭиСО"

    @staticmethod
    def _get_search_dirs():
        """Возвращает список папок для поиска в зависимости от ОС."""
        candidates = []

        if platform == "android":
            candidates.append(App.get_running_app().user_data_dir)
            candidates.append("/storage/emulated/0/Download")
            candidates.append("/storage/emulated/0/Documents")
            candidates.append("/storage/emulated/0")
        else:
            up = os.environ.get("USERPROFILE")
            if up:
                candidates.append(os.path.join(up, "Downloads"))
            home = os.path.expanduser("~")
            candidates.append(os.path.join(home, "Downloads"))
            candidates.append(os.path.join(home, "Загрузки"))

        valid_dirs = [d for d in candidates if os.path.isdir(d)]
        return valid_dirs if valid_dirs else ["/"]

    def build(self):
        self.constants_dict = load_config()

        root = BoxLayout(orientation='vertical', padding=10, spacing=8)

        root.add_widget(Label(
            text="Формирование ежедневной сводки ЦТЭиСО",
            size_hint_y=None, height=30, bold=True, font_size=18
        ))

        # Поле 1: site_not_air
        box1 = BoxLayout(size_hint_y=None, height=40, spacing=4)
        self.inp_om = TextInput(hint_text="Путь к site_not_air_*.xlsx", multiline=False)
        btn1 = Button(text="...", size_hint_x=None, width=50)
        btn1.bind(on_release=lambda x: self.pick_file(self.inp_om, "site_not_air"))
        box1.add_widget(self.inp_om)
        box1.add_widget(btn1)
        root.add_widget(box1)

        # Поле 2: Список аварий
        box2 = BoxLayout(size_hint_y=None, height=40, spacing=4)
        self.inp_alarms = TextInput(hint_text="Путь к Список аварий*.xlsx", multiline=False)
        btn2 = Button(text="...", size_hint_x=None, width=50)
        btn2.bind(on_release=lambda x: self.pick_file(self.inp_alarms, "Список аварий"))
        box2.add_widget(self.inp_alarms)
        box2.add_widget(btn2)
        root.add_widget(box2)

        # Поле 3: СЦ (Spinner)
        box3 = BoxLayout(size_hint_y=None, height=40, spacing=4)
        box3.add_widget(Label(text="СЦ:", size_hint_x=None, width=50))
        self.spinner_sc = Spinner(
            text="Сервисный Центр Феодосия",
            values=("Сервисный Центр Феодосия", "Сервисный Центр Симферополь"),
            size_hint=(1, None), height=40
        )
        box3.add_widget(self.spinner_sc)
        root.add_widget(box3)

        # Поле 4: Константы
        box4 = BoxLayout(size_hint_y=None, height=100, spacing=4, orientation='vertical')
        box4.add_widget(
            Label(text="Константы (Правила БС в формате JSON):", size_hint_y=None, height=25, halign="left"))

        self.inp_constants = TextInput(
            text=json.dumps(self.constants_dict, ensure_ascii=False, indent=2),
            multiline=True,
            font_size=12
        )
        box4.add_widget(self.inp_constants)
        root.add_widget(box4)

        # Кнопка запуска
        btn_gen = Button(
            text="Сформировать отчёт",
            size_hint_y=None, height=45,
            background_color=(0.2, 0.6, 0.8, 1)
        )
        btn_gen.bind(on_release=self.generate)
        root.add_widget(btn_gen)

        # Вывод
        self.out_text = TextInput(readonly=True, multiline=True, font_size=14)
        self.out_text.background_color = (0.95, 0.95, 0.95, 1)
        scroll = ScrollView()
        scroll.add_widget(self.out_text)
        root.add_widget(scroll)

        return root

    def pick_file(self, target_input, prefix):
        search_dirs = self._get_search_dirs()

        search_box = BoxLayout(size_hint_y=None, height=40, spacing=4)
        search_box.add_widget(Label(text="Поиск:", size_hint_x=None, width=60))
        search_input = TextInput(multiline=False, hint_text="Введите часть имени...")
        search_box.add_widget(search_input)

        fc_container = BoxLayout()

        row = BoxLayout(size_hint_y=None, height=48, spacing=4)
        btn_ok = Button(text="Выбрать")
        btn_cn = Button(text="Отмена")
        row.add_widget(btn_ok)
        row.add_widget(btn_cn)

        content = BoxLayout(orientation="vertical", spacing=4)
        content.add_widget(search_box)
        content.add_widget(fc_container)
        content.add_widget(row)

        popup = Popup(title=f"Выбор файла ({prefix}...)", content=content, size_hint=(0.95, 0.95))

        def is_valid_file(filename):
            name_lower = filename.lower()
            if prefix == "site_not_air":
                return "site_not_air" in name_lower or "not_air" in name_lower
            elif prefix == "Список аварий":
                return "авари" in name_lower
            return True

        def update_file_list(search_text=""):
            fc_container.clear_widgets()
            list_view = ScrollView()
            list_layout = BoxLayout(orientation="vertical", size_hint_y=None)
            list_layout.bind(minimum_height=list_layout.setter('height'))

            found_files = []

            for d in search_dirs:
                try:
                    for f in os.listdir(d):
                        full_path = os.path.join(d, f)
                        if os.path.isfile(full_path) and is_valid_file(f):
                            if search_text.lower() in f.lower():
                                found_files.append(full_path)
                except Exception:
                    pass

            found_files = list(set(found_files))

            if found_files:
                for file_path in found_files:
                    file_name = os.path.basename(file_path)
                    btn = Button(text=file_name, size_hint_y=None, height=50, halign="left")
                    btn.text_size = (btn.width, None)
                    btn.bind(width=lambda inst, val: setattr(inst, 'text_size', (val, None)))
                    btn.bind(on_release=lambda x, path=file_path: select_found_file(path))
                    list_layout.add_widget(btn)
            else:
                list_layout.add_widget(Label(text="Файлы не найдены", size_hint_y=None, height=40))

            list_view.add_widget(list_layout)
            fc_container.add_widget(list_view)

        def select_found_file(full_path):
            target_input.text = full_path
            if target_input == self.inp_om:
                try:
                    scs = get_unique_scs(full_path)
                    self.spinner_sc.values = scs
                    self.spinner_sc.text = scs[0]
                except Exception as e:
                    print(f"Не удалось прочитать СЦ: {e}")
            popup.dismiss()

        search_input.bind(text=lambda instance, text: update_file_list(text))
        update_file_list()

        def _ok(*_):
            pass

        btn_ok.bind(on_release=_ok)
        btn_cn.bind(on_release=lambda *_: popup.dismiss())
        popup.open()

    def generate(self, instance):
        try:
            raw_text = self.inp_constants.text.strip()
            try:
                constants = json.loads(raw_text)
            except json.JSONDecodeError:
                raise ValueError("Ошибка в JSON формате констант! Проверьте кавычки и запятые.")

            save_config(constants)

            self.out_text.text = build_report(
                self.inp_om.text,
                self.inp_alarms.text,
                self.spinner_sc.text,
                constants
            )
        except Exception as e:
            self.out_text.text = f"Ошибка при формировании отчёта:\n{e}\n\n{traceback.format_exc()}"


if __name__ == "__main__":
    ReportApp().run()